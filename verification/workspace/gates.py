"""L4 chat workspace gates (``deploy/workspace``): one implementation for every example.

Docling OCR, Japanese knowledge bases, the Chat UI in a browser, serving
beside OCR load, Word / Excel / PowerPoint and download links. An example's
``verification.py`` loads this file by path (its own module is named
``verification``), builds :class:`WorkspaceGates` with its values and adds
``gates()`` to its gate list. The fixtures are generated Japanese documents:
an image-only "scanned" PDF rendered with a pinned Noto Sans JP, a text-layer
PDF printed by Chromium, and plain-text distractors that share the scanned
document's vocabulary. Thresholds come from the workspace profile's ``gates``.
"""

from __future__ import annotations

import hashlib
import json
import os
import re
import subprocess
import sys
import threading
import time
import urllib.error
import urllib.request
from collections.abc import Callable
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
BROWSER_SCRIPT = Path(__file__).with_name("browser.mjs")
sys.path.append(str(ROOT / "deploy" / "workspace"))
from kairyu_workspace import (  # noqa: E402
    COMPOSE_FILE,
    OPEN_WEBUI_IMAGE,
    Workspace,
    chat_ui_port,
    compose_env,
)
from kairyu_workspace.openwebui import sign_in  # noqa: E402
from kairyu_workspace.spec import DOCLING_URL  # noqa: E402

KNOWLEDGE_NAME = "kairyu-knowledge-gate"
KNOWLEDGE_SCAN = "kairyu-gate-scan.pdf"
KNOWLEDGE_KEY_SENTENCE = (
    "保存条件：試料KR-7319は摂氏マイナス四十二度の冷凍庫で保管し、解凍は一回までとする。"
)
# Fragments without glyph twins: OCR may read the kanji "一" as the long-vowel "ー".
KNOWLEDGE_KEY_FRAGMENTS = ("KR-7319", "マイナス四十二度", "冷凍庫で保管")
KNOWLEDGE_QUERY = "KR-7319という試料は、どのような温度で保存するよう決められていますか。"
_JA_SUBJECTS = (
    "第一分析室",
    "第二分析室",
    "品質管理課",
    "設備保全班",
    "試料受付窓口",
    "安全衛生委員会",
    "製剤開発グループ",
    "物性評価チーム",
)
_JA_TEMPLATES = (
    "{subject}では、{month}月の点検で{device}の校正記録を{count}件確認した。",
    "{subject}は、試料{code}の受入時に外観と重量を記録し、異常がないことを確かめた。",
    "{device}の稼働時間は月平均{count}時間で、前年同月より{delta}時間増えた。",
    "{subject}の報告によると、試料{code}は摂氏{temp}度の保管庫に移された。",
    "手順書の改訂案では、{device}の清掃を{count}日ごとに行うことが提案されている。",
    "{month}月の会議で、{subject}は廃液の分別方法について{count}項目の見直しを求めた。",
    "試料{code}の分析結果は、規格の範囲内であったため再試験は行わない。",
    "{subject}は、{device}の消耗品を{count}個発注し、納期を{month}月末とした。",
)
_JA_DEVICES = ("液体クロマトグラフ", "分光光度計", "恒温槽", "遠心分離機", "電子天秤", "乾燥機")
_JA_TEMPS = ("四", "五", "マイナス二十", "マイナス八十", "二十五", "マイナス四十")


def _normalized(text: str) -> str:
    """NFKC, no whitespace, every dash-like character as ``-`` (OCR spacing varies)."""
    import unicodedata

    text = re.sub(r"\s+", "", unicodedata.normalize("NFKC", text))
    text = re.sub("[\u2010-\u2015\u2212]", "-", text)
    # OCR may read the sample code's hyphen as a long-vowel mark or a kanji one.
    return re.sub("KR[-\u30fc\u4e00]?(?=\\d)", "KR-", text)


def _has_key(text: str) -> bool:
    flat = _normalized(text)
    return all(_normalized(fragment) in flat for fragment in KNOWLEDGE_KEY_FRAGMENTS)


def knowledge_lines(page: int, *, key: bool, namespace: str) -> list[str]:
    """Deterministic Japanese report lines for one page; ``key`` plants the key sentence."""
    lines = [f"技術報告書 {namespace} 第{page}頁"]
    for index in range(16):
        digest = hashlib.sha256(f"{namespace}:{page}:{index}".encode()).digest()
        code = f"KR-{7300 + digest[1] % 90 + (1 if 7300 + digest[1] % 90 >= 7319 else 0)}"
        lines.append(
            _JA_TEMPLATES[digest[0] % len(_JA_TEMPLATES)].format(
                subject=_JA_SUBJECTS[digest[2] % len(_JA_SUBJECTS)],
                device=_JA_DEVICES[digest[3] % len(_JA_DEVICES)],
                month=1 + digest[4] % 12,
                count=2 + digest[5] % 30,
                delta=1 + digest[6] % 9,
                temp=_JA_TEMPS[digest[7] % len(_JA_TEMPS)],
                code=code,
            )
        )
        if key and index == 7:
            lines.append(KNOWLEDGE_KEY_SENTENCE)
    return lines


def distractor_documents(count: int, *, namespace: str) -> list[tuple[str, str]]:
    return [
        (
            f"kairyu-gate-note-{index:02d}.md",
            "\n".join(knowledge_lines(100 + index, key=False, namespace=namespace)) + "\n",
        )
        for index in range(1, count + 1)
    ]


_LOADER_PROGRAM = r"""
import json, os, sys, time
sys.path.insert(0, '/app/backend')
os.makedirs(os.environ['DATA_DIR'], exist_ok=True)
from open_webui.retrieval.loaders.main import DoclingLoader
params = json.loads(os.environ['DOCLING_PARAMS'])
report = {}
for name in sys.argv[1:]:
    started = time.monotonic()
    documents = DoclingLoader(url=os.environ['DOCLING_SERVER_URL'], file_path='/fixtures/' + name,
                              mime_type='application/pdf', params=params).load()
    report[name] = {'seconds': time.monotonic() - started,
                    'pages': [document.page_content for document in documents]}
print('LOADER-REPORT ' + json.dumps(report, ensure_ascii=False))
"""


OFFICE_CASES = {
    "docx": {
        "fixture": "kairyu-office-report.docx",
        "result": "kairyu-office-report-revised.docx",
        "question": "添付の報告書によると、納品期限はいつですか。日付だけ答えてください。",
        "answer": r"(2026\s*年\s*)?12\s*月\s*18\s*日|2026[-/.]12[-/.]18|12/18",
        "edit": (
            "添付の kairyu-office-report.docx を編集してください。本文の「東京本社」をすべて"
            "「大阪支社」に置き換え、文書の最後に「改訂日: 2026年10月22日」という段落を追加し、"
            "/mnt/uploads/kairyu-office-report-revised.docx として保存してください。"
            "表やほかの本文は変えないでください。"
        ),
    },
    "xlsx": {
        "fixture": "kairyu-office-sales.xlsx",
        "result": "kairyu-office-sales-total.xlsx",
        "question": "添付の売上表で、2026年9月の売上はいくらですか。数値だけ答えてください。",
        "answer": r"\b161\b",
        "edit": (
            "添付の kairyu-office-sales.xlsx の「売上」シートの最終行の下に、A列が「合計」、"
            "B列がB2からB7の合計（SUM 関数）の行を追加し、ほかのセルは変えずに "
            "/mnt/uploads/kairyu-office-sales-total.xlsx として保存してください。"
        ),
    },
    "pptx": {
        "fixture": "kairyu-office-review.pptx",
        "result": "kairyu-office-review-final.pptx",
        "question": "添付の資料によると、新規顧客は何社増えましたか。数だけ答えてください。",
        "answer": r"\b42\b|四十二",
        "edit": (
            "添付の kairyu-office-review.pptx の最後に、タイトルが「まとめ」で、本文に"
            "箇条書きが3行あるスライドを1枚追加し、既存のスライドは変えずに "
            "/mnt/uploads/kairyu-office-review-final.pptx として保存してください。"
        ),
    },
}

_OFFICE_FIXTURES = r"""
import docx, openpyxl, pptx, sys
from pptx.util import Pt
out = sys.argv[1]
d = docx.Document()
d.add_heading('出張報告書', 1)
d.add_paragraph('出張先は東京本社で、期間は2026年11月4日から6日までとする。')
d.add_paragraph('納品期限は2026年12月18日で、担当は品質管理課の佐藤とする。')
d.add_paragraph('東京本社の会議室は第3会議室を予約済みである。')
table = d.add_table(rows=1, cols=3)
for cell, text in zip(table.rows[0].cells, ('試料', '数量', '保管場所')):
    cell.text = text
for row in (('試料A', '12', '東京本社'), ('試料B', '8', '名古屋'), ('試料C', '5', '福岡')):
    for cell, text in zip(table.add_row().cells, row):
        cell.text = text
d.save(f'{out}/kairyu-office-report.docx')
book = openpyxl.Workbook()
sheet = book.active
sheet.title = '売上'
sheet.append(['月', '売上（万円）'])
for month, value in zip(range(4, 10), (120, 135, 128, 142, 150, 161)):
    sheet.append([f'2026年{month}月', value])
book.save(f'{out}/kairyu-office-sales.xlsx')
deck = pptx.Presentation()
for title, lines in (('第3四半期の振り返り', ['2026年7月から9月']),
                     ('主な成果', ['新規顧客は42社増えた', '解約率は1.8%に下がった']),
                     ('課題', ['納期遅延が3件あった', '問い合わせ対応の待ち時間'])):
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    slide.shapes.title.text = title
    body = slide.placeholders[1].text_frame
    body.text = lines[0]
    for line in lines[1:]:
        body.add_paragraph().text = line
deck.save(f'{out}/kairyu-office-review.pptx')
print('fixtures ok')
"""

_OFFICE_CHECK = r"""
import docx, json, openpyxl, os, pptx, sys
out = sys.argv[1]
report = {}
def check(name, function):
    path = os.path.join(out, name)
    if not os.path.exists(path):
        report[name] = {'exists': False, 'passed': False}
        return
    try:
        report[name] = {'exists': True, **function(path)}
    except Exception as error:
        report[name] = {'exists': True, 'passed': False, 'error': repr(error)}
def revised_docx(path):
    d = docx.Document(path)
    text = '\n'.join(p.text for p in d.paragraphs)
    rows = [[c.text for c in r.cells] for r in d.tables[0].rows] if d.tables else []
    last = next((p.text for p in reversed(d.paragraphs) if p.text.strip()), '')
    result = {
        'replaced': '大阪支社' in text and '東京本社' not in text,
        'appended': '改訂日' in last and '2026年10月22日' in last,
        'kept_sentence': '納品期限は2026年12月18日' in text,
        'kept_table_rows': len(rows) == 4 and rows[0] == ['試料', '数量', '保管場所'],
    }
    return {**result, 'passed': all(result.values())}
def total_xlsx(path):
    sheet = openpyxl.load_workbook(path)['売上']
    rows = list(sheet.iter_rows(values_only=True))
    kept = rows[:7] == [('月', '売上（万円）')] + [
        (f'2026年{m}月', v) for m, v in zip(range(4, 10), (120, 135, 128, 142, 150, 161))]
    total = rows[7] if len(rows) > 7 else (None, None)
    value = str(total[1]).replace(' ', '').upper()
    result = {
        'kept_rows': kept,
        'total_label': total[0] == '合計',
        'total_sum': value in ('=SUM(B2:B7)', '836', '836.0'),
    }
    return {**result, 'total_cell': str(total[1]), 'passed': all(result.values())}
def final_pptx(path):
    deck = pptx.Presentation(path)
    titles = [s.shapes.title.text if s.shapes.title else '' for s in deck.slides]
    last = deck.slides[-1]
    lines = [p.text for shape in last.shapes if shape.has_text_frame and shape != last.shapes.title
             for p in shape.text_frame.paragraphs if p.text.strip()]
    result = {
        'kept_titles': titles[:3] == ['第3四半期の振り返り', '主な成果', '課題'],
        'added_slide': len(titles) == 4 and titles[3] == 'まとめ',
        'three_bullets': len(lines) >= 3,
    }
    return {**result, 'titles': titles, 'passed': all(result.values())}
check('kairyu-office-report-revised.docx', revised_docx)
check('kairyu-office-sales-total.xlsx', total_xlsx)
check('kairyu-office-review-final.pptx', final_pptx)
print('OFFICE-CHECK ' + json.dumps(report, ensure_ascii=False))
"""


# The owner's own request (2026-10-22, on a phone) that produced a file the
# Chat UI could not hand over.
DOWNLOAD_REQUEST = "一年の事業計画のフォーマットをエクセルで作って"
_DOWNLOAD_CHECK = r"""
import hashlib, json, openpyxl, os, sys
out = sys.argv[1]
report = {}
for name in ("desktop.xlsx", "desktop-after-reload.xlsx", "phone.xlsx"):
    path = os.path.join(out, name)
    if not os.path.exists(path):
        report[name] = {"exists": False, "passed": False}
        continue
    try:
        book = openpyxl.load_workbook(path)
        cells = sum(
            1 for sheet in book.worksheets for row in sheet.iter_rows(values_only=True)
            for value in row if value not in (None, "")
        )
        report[name] = {
            "exists": True,
            "sheets": book.sheetnames,
            "filled_cells": cells,
            "sha256": hashlib.sha256(open(path, "rb").read()).hexdigest(),
            "passed": bool(book.sheetnames) and cells >= 10,
        }
    except Exception as error:
        report[name] = {"exists": True, "passed": False, "error": repr(error)}
first, again = report["desktop.xlsx"], report["desktop-after-reload.xlsx"]
digest = first.get("sha256")
report["same_after_reload"] = bool(digest) and digest == again.get("sha256")
print("DOWNLOAD-CHECK " + json.dumps(report, ensure_ascii=False))
"""


class DoclingCpu:
    """Docling's own CPU use, sampled once a second from its cgroup counter.

    This is the OCR load itself: Open WebUI's upload, queue and embedding time
    does not count, only CPU time the Docling container spends.
    """

    def __init__(self, container: str, interval_s: float = 1.0) -> None:
        import threading

        identity = subprocess.run(
            ["docker", "inspect", "--format", "{{.Id}}", container],
            capture_output=True,
            text=True,
            check=True,
        ).stdout.strip()
        candidates = (
            Path(f"/sys/fs/cgroup/system.slice/docker-{identity}.scope/cpu.stat"),
            Path(f"/sys/fs/cgroup/docker/{identity}/cpu.stat"),
        )
        self.path = next((path for path in candidates if path.exists()), None)
        if self.path is None:
            raise RuntimeError(f"no cgroup CPU counter for {container}")
        self.interval_s = interval_s
        # (start, end, cores used) per sample, monotonic seconds.
        self.samples: list[tuple[float, float, float]] = []
        self.error: str | None = None
        self._stop = threading.Event()
        self._thread = threading.Thread(target=self._run, daemon=True)

    def _usage_s(self) -> float:
        for line in self.path.read_text().splitlines():
            key, value = line.split()
            if key == "usage_usec":
                return int(value) / 1e6
        raise RuntimeError(f"{self.path} has no usage_usec")

    def _run(self) -> None:
        # A failed read ends sampling; ``error`` fails the gate and the time
        # after the last sample counts as unobserved, never as busy.
        try:
            last_time, last_usage = time.monotonic(), self._usage_s()
            while not self._stop.wait(self.interval_s):
                now, usage = time.monotonic(), self._usage_s()
                self.samples.append((last_time, now, (usage - last_usage) / (now - last_time)))
                last_time, last_usage = now, usage
        except Exception as error:  # noqa: BLE001 - reported through ``error``
            self.error = repr(error)

    def __enter__(self) -> DoclingCpu:
        self._thread.start()
        return self

    def __exit__(self, *_: object) -> None:
        self._stop.set()
        self._thread.join()

    def wait_busy(self, cores: float, seconds: int, timeout_s: float) -> bool:
        """True once the last ``seconds`` samples each used at least ``cores``."""
        deadline = time.monotonic() + timeout_s
        while time.monotonic() < deadline and self.error is None:
            recent = self.samples[-seconds:]
            if len(recent) == seconds and all(used >= cores for *_, used in recent):
                return True
            time.sleep(self.interval_s / 2)
        return False

    def wait_idle(self, cores: float, seconds: int, timeout_s: float) -> bool:
        """True once the last ``seconds`` samples each used less than ``cores``."""
        deadline = time.monotonic() + timeout_s
        while time.monotonic() < deadline and self.error is None:
            recent = self.samples[-seconds:]
            if len(recent) == seconds and all(used < cores for *_, used in recent):
                return True
            time.sleep(self.interval_s / 2)
        return False

    def wait_past(self, moment: float) -> None:
        """Wait (up to three intervals) for a sample that ends after ``moment``."""
        deadline = time.monotonic() + 3 * self.interval_s
        while time.monotonic() < deadline and self.error is None:
            if self.samples and self.samples[-1][1] >= moment:
                return
            time.sleep(self.interval_s / 4)

    def summary(self, started: float, ended: float, cores: float) -> dict:
        """Busy and observed shares of [started, ended], weighted by time.

        Seconds without a sample (a stopped sampler) are unobserved and so not
        busy: the busy share's denominator is the whole window.
        """
        window = max(ended - started, 0.0)
        observed = busy = core_seconds = 0.0
        for first, last, used in self.samples:
            overlap = max(0.0, min(last, ended) - max(first, started))
            observed += overlap
            core_seconds += overlap * used
            if used >= cores:
                busy += overlap
        return {
            "busy_share": busy / window if window else 0.0,
            "observed_share": observed / window if window else 0.0,
            "mean_cores": core_seconds / observed if observed else 0.0,
        }


class WorkspaceGates:
    """The L4 gates against one example's running workspace.

    ``serving(run_dir, concurrency)`` runs the example's own serving rows and
    ``serving_tok_s(run_dir, level)`` reads one row's output tokens per second;
    ``knowledge_serving`` holds the example's reference for those rows.
    """

    def __init__(
        self,
        workspace: Workspace,
        *,
        project: str,
        results_root: Path,
        work_dir: Path,
        serving: Callable[[Path, list[int]], int],
        serving_tok_s: Callable[[Path, int], float | None],
        knowledge_serving: dict,
    ) -> None:
        self.workspace = workspace
        self.limits = workspace.profile["gates"]
        self.project = project
        self.work_dir = work_dir
        self.serving = serving
        self.serving_tok_s = serving_tok_s
        self.reference = knowledge_serving
        self.ui_url = f"http://127.0.0.1:{chat_ui_port(workspace)}"
        # Every file a gate uploads is recorded under its knowledge base, so
        # cleanup deletes exactly those (failed uploads included), never a
        # user's file.
        self.uploads = results_root / "knowledge-gate-uploads.json"
        self._uploads_lock = threading.Lock()

    def gates(self) -> dict[str, Callable[[Path], int]]:
        return {
            "docling": self.docling,
            "knowledge": self.knowledge,
            "chat-ui": self.chat_ui,
            "knowledge-serving": self.knowledge_serving,
            "office": self.office,
            "file-download": self.file_download,
        }

    def runtime_evidence(self) -> dict:
        """The running Chat UI and Docling must be L4's pinned images."""
        pins = {"chat-ui": OPEN_WEBUI_IMAGE, "docling": self.workspace.profile["docling"]["image"]}
        evidence = {}
        for service, pin in pins.items():
            name = f"{self.project}-{service}-1"
            container = json.loads(subprocess.check_output(["docker", "inspect", name]))[0]
            if container["Config"]["Image"] != pin:
                raise ValueError(f"running {service} image differs from the L4 pin")
            evidence[service] = {
                "image": pin,
                "image_id": container["Image"],
                "started_at": container["State"]["StartedAt"],
            }
        return evidence

    # --- fixtures ----------------------------------------------------------------------

    def knowledge_font(self) -> Path:
        """The pinned Noto Sans JP (OFL) used to render the Japanese fixtures."""
        target = self.work_dir / "NotoSansJP-Regular.otf"
        if not target.exists():
            target.parent.mkdir(parents=True, exist_ok=True)
            with urllib.request.urlopen(self.limits["font_url"], timeout=120) as response:
                target.write_bytes(response.read())
        digest = hashlib.sha256(target.read_bytes()).hexdigest()
        if digest != self.limits["font_sha256"]:
            target.unlink()
            raise RuntimeError(f"Noto Sans JP digest {digest} differs from the pinned font")
        return target

    def scan_pdf(self, path: Path, pages: int, key_pages: set[int], *, namespace: str) -> None:
        """An image-only A4 PDF at 150 dpi: there is no text layer, only pixels to OCR."""
        from PIL import Image, ImageDraw, ImageFont

        font = ImageFont.truetype(str(self.knowledge_font()), 28)
        width, height, margin, line_height = 1240, 1754, 110, 46
        images = []
        for page in range(1, pages + 1):
            image = Image.new("L", (width, height), 255)
            draw = ImageDraw.Draw(image)
            y = margin
            for line in knowledge_lines(page, key=page in key_pages, namespace=namespace):
                row = ""
                for character in line:
                    if draw.textlength(row + character, font=font) > width - 2 * margin:
                        draw.text((margin, y), row, font=font, fill=0)
                        y, row = y + line_height, ""
                    row += character
                draw.text((margin, y), row, font=font, fill=0)
                y += line_height + 10
            images.append(image)
        path.parent.mkdir(parents=True, exist_ok=True)
        images[0].save(path, save_all=True, append_images=images[1:], resolution=150)

    def browser_image(self) -> str:
        """Build the pinned Playwright image (deploy/compose/Dockerfile.webui-browser)."""
        image = self.limits["browser_image"]
        build = subprocess.run(
            [
                "docker",
                "build",
                "--file",
                str(ROOT / "deploy/compose/Dockerfile.webui-browser"),
                "--tag",
                image,
                str(ROOT),
            ],
            capture_output=True,
            text=True,
        )
        if build.returncode:
            raise RuntimeError(f"building {image} failed: {build.stderr[-2000:]}")
        return image

    def text_pdf(self, path: Path, lines: list[str]) -> None:
        """A text-layer PDF of ``lines``, printed by Chromium with the pinned font embedded."""
        import html

        work = path.parent
        font = self.knowledge_font()
        (work / font.name).write_bytes(font.read_bytes())
        body = "".join(f"<p>{html.escape(line)}</p>" for line in lines)
        (work / "text.html").write_text(
            "<!doctype html><meta charset='utf-8'><style>"
            f"@font-face{{font-family:jp;src:url('{font.name}')}}"
            "body{font-family:jp;font-size:14pt;margin:2cm}</style>" + body,
            encoding="utf-8",
        )
        subprocess.run(
            [
                "docker",
                "run",
                "--rm",
                "--network",
                "none",
                "--user",
                f"{os.getuid()}:{os.getgid()}",
                "--env",
                "HOME=/tmp",
                "--volume",
                f"{work}:/fixtures",
                "--entrypoint",
                "/ms-playwright/chromium-1223/chrome-linux64/chrome",
                self.browser_image(),
                "--headless",
                "--no-sandbox",
                "--disable-gpu",
                "--disable-dev-shm-usage",
                "--disable-crash-reporter",
                "--no-pdf-header-footer",
                f"--print-to-pdf=/fixtures/{path.name}",
                "file:///fixtures/text.html",
            ],
            check=True,
            capture_output=True,
        )

    def _compose_service(self, name: str) -> dict:
        """A service of the L4 compose file with this workspace's values filled in."""
        import yaml

        values = compose_env(self.workspace, data=Path("/unused"), models=Path("/unused"))
        pattern = re.compile(r"\$\{(\w+)(?::([?-])([^}]*))?\}")

        def fill(value):
            if isinstance(value, dict):
                return {key: fill(item) for key, item in value.items()}
            if isinstance(value, list):
                return [fill(item) for item in value]
            if not isinstance(value, str):
                return value
            return pattern.sub(
                lambda match: (
                    values.get(match[1])
                    or os.environ.get(match[1])
                    or (match[3] if match[2] == "-" else "")
                ),
                value,
            )

        services = yaml.safe_load(COMPOSE_FILE.read_text(encoding="utf-8"))["services"]
        return fill(services[name])

    def _docling_python(self, program: str, directory: Path, *arguments: str) -> str:
        """Run ``program`` offline in the pinned Docling image with ``directory`` at /work."""
        run = subprocess.run(
            [
                "docker",
                "run",
                "--rm",
                "--network",
                "none",
                "--user",
                f"{os.getuid()}:{os.getgid()}",
                "--volume",
                f"{directory}:/work",
                "--entrypoint",
                "python3",
                self.workspace.profile["docling"]["image"],
                "-c",
                program,
                *arguments,
            ],
            capture_output=True,
            text=True,
            timeout=600,
        )
        if run.returncode:
            raise RuntimeError(f"Docling-image program failed: {run.stderr[-2000:]}")
        return run.stdout

    # --- Open WebUI ----------------------------------------------------------------------

    def _token(self) -> str:
        return sign_in(
            self.ui_url,
            self.workspace.auth,
            os.environ.get("WEBUI_ADMIN_EMAIL", ""),
            os.environ.get("WEBUI_ADMIN_PASSWORD", ""),
        )

    def _request(self, method: str, path: str, token: str, *, payload=None, timeout_s=600.0):
        request = urllib.request.Request(
            f"{self.ui_url}{path}",
            data=json.dumps(payload).encode() if payload is not None else None,
            headers={"Authorization": f"Bearer {token}", "Content-Type": "application/json"},
            method=method,
        )
        with urllib.request.urlopen(request, timeout=timeout_s) as response:
            return json.loads(response.read() or b"null")

    def _upload(self, token: str, name: str, data: bytes, content_type: str) -> dict:
        boundary = f"kairyu{hashlib.sha256(data).hexdigest()[:24]}"
        body = (
            (
                f"--{boundary}\r\n"
                f'Content-Disposition: form-data; name="file"; filename="{name}"\r\n'
                f"Content-Type: {content_type}\r\n\r\n"
            ).encode()
            + data
            + f"\r\n--{boundary}--\r\n".encode()
        )
        request = urllib.request.Request(
            f"{self.ui_url}/api/v1/files/?process=true&process_in_background=true",
            data=body,
            headers={
                "Authorization": f"Bearer {token}",
                "Content-Type": f"multipart/form-data; boundary={boundary}",
            },
            method="POST",
        )
        with urllib.request.urlopen(request, timeout=600) as response:
            return json.loads(response.read())

    def _knowledge_by_name(self, token: str, name: str) -> dict | None:
        listed = self._request("GET", "/api/v1/knowledge/", token)
        rows = listed.get("items", listed) if isinstance(listed, dict) else listed
        return next((row for row in rows or [] if row.get("name") == name), None)

    def _knowledge_files(self, token: str, knowledge_id: str) -> list[dict]:
        # The knowledge detail's ``files`` is empty in v0.12.3; the list has its own route.
        listed = self._request("GET", f"/api/v1/knowledge/{knowledge_id}/files", token)
        return listed.get("items", []) if isinstance(listed, dict) else listed or []

    def _gate_uploads(self) -> dict[str, list[str]]:
        return json.loads(self.uploads.read_text()) if self.uploads.exists() else {}

    def _record_upload(self, knowledge_id: str, file_id: str) -> None:
        with self._uploads_lock:
            uploads = self._gate_uploads()
            uploads.setdefault(knowledge_id, []).append(file_id)
            self.uploads.parent.mkdir(parents=True, exist_ok=True)
            self.uploads.write_text(json.dumps(uploads, indent=2) + "\n")

    def _delete_knowledge(self, token: str, knowledge_id: str) -> None:
        """Delete a gate knowledge base and the files the gates recorded for it."""
        try:
            self._request("DELETE", f"/api/v1/knowledge/{knowledge_id}/delete", token)
        except urllib.error.HTTPError as error:
            if error.code != 404:
                raise
        with self._uploads_lock:
            uploads = self._gate_uploads()
            for file_id in uploads.pop(knowledge_id, []):
                try:
                    self._request("DELETE", f"/api/v1/files/{file_id}", token)
                except urllib.error.HTTPError as error:
                    if error.code != 404:
                        raise
            if self.uploads.exists():
                self.uploads.write_text(json.dumps(uploads, indent=2) + "\n")

    def _ingest(
        self, token: str, knowledge_id: str, name: str, data: bytes, content_type: str
    ) -> tuple[str, float]:
        """Upload, wait for Open WebUI's extraction, add to the knowledge base; seconds taken."""
        started = time.monotonic()
        file_id = self._upload(token, name, data, content_type)["id"]
        self._record_upload(knowledge_id, file_id)
        deadline = started + 2 * float(self.limits["ingestion_seconds"])
        while True:
            status = self._request("GET", f"/api/v1/files/{file_id}/process/status", token).get(
                "status"
            )
            if status == "completed":
                break
            if status == "failed" or time.monotonic() > deadline:
                detail = self._request("GET", f"/api/v1/files/{file_id}", token)
                raise RuntimeError(
                    f"{name}: processing {status}: {json.dumps(detail.get('data'))[:500]}"
                )
            time.sleep(1)
        self._request(
            "POST",
            f"/api/v1/knowledge/{knowledge_id}/file/add",
            token,
            payload={"file_id": file_id},
            timeout_s=1800,
        )
        return file_id, time.monotonic() - started

    def _kairyu_api_counts(self) -> dict[str, int]:
        logs = subprocess.run(
            ["docker", "logs", f"{self.project}-kairyu-1"], capture_output=True, text=True
        )
        text = logs.stdout + logs.stderr
        return {
            path: len(re.findall(rf'"path": "{re.escape(path)}"', text))
            for path in ("/v1/responses", "/v1/chat/completions")
        }

    def _browser(
        self, run_dir: Path, environment: dict[str, str], volumes: list[str]
    ) -> subprocess.CompletedProcess:
        """Run browser.mjs against the Chat UI in the pinned Playwright image."""
        command = ["docker", "run", "--rm", "--init", "--network", "host"]
        environment = {
            "WEBUI_SMOKE_BASE_URL": self.ui_url,
            "WEBUI_SMOKE_MODELS": json.dumps(list(self.workspace.models)),
            **environment,
        }
        for key, value in environment.items():
            command += ["--env", f"{key}={value}"]
        for volume in (f"{BROWSER_SCRIPT}:/work/webui_browser_smoke.mjs:ro", *volumes):
            command += ["--volume", volume]
        command += [self.browser_image(), "node", "/work/webui_browser_smoke.mjs"]
        browser = subprocess.run(command, capture_output=True, text=True, timeout=3600)
        (run_dir / "browser.log").write_text(browser.stdout + browser.stderr)
        return browser

    # --- gates ---------------------------------------------------------------------------

    def docling(self, run_dir: Path) -> int:
        """Docling alone, offline: Japanese image-only and text-layer PDFs, Open WebUI's loader."""
        fixtures = run_dir / "fixtures"
        pages = int(self.limits["probe_pages"])
        self.scan_pdf(
            fixtures / "scan.pdf", pages, set(range(1, pages + 1)), namespace=run_dir.name
        )
        text_lines = knowledge_lines(1, key=True, namespace=f"{run_dir.name}-text")
        self.text_pdf(fixtures / "text.pdf", text_lines)
        service = self._compose_service("docling")
        image = self.workspace.profile["docling"]["image"]
        params = json.dumps(self.workspace.profile["docling"]["params"], separators=(",", ":"))
        network = f"kairyu-docling-probe-{run_dir.name.lower()}"
        container = f"{network}-docling"
        environment = [
            item
            for key, value in service["environment"].items()
            for item in ("--env", f"{key}={value}")
        ]
        report: dict = {"docling_image": image, "docling_params": params}
        subprocess.run(
            ["docker", "network", "create", "--internal", network], check=True, capture_output=True
        )
        try:
            subprocess.run(
                [
                    "docker",
                    "run",
                    "--detach",
                    "--name",
                    container,
                    "--network",
                    network,
                    "--network-alias",
                    "docling",
                    "--cpus",
                    str(service["cpus"]),
                    *environment,
                    image,
                ],
                check=True,
                capture_output=True,
            )
            deadline = time.monotonic() + 300
            while subprocess.run(
                [
                    "docker",
                    "exec",
                    container,
                    "python3",
                    "-c",
                    "import urllib.request; "
                    "urllib.request.urlopen('http://127.0.0.1:5001/health', timeout=2)",
                ],
                capture_output=True,
            ).returncode:
                if time.monotonic() > deadline:
                    raise RuntimeError("Docling did not become healthy")
                time.sleep(2)
            loaded = subprocess.run(
                [
                    "docker",
                    "run",
                    "--rm",
                    "--network",
                    network,
                    "--volume",
                    f"{fixtures}:/fixtures:ro",
                    "--env",
                    "DATA_DIR=/tmp/open-webui",
                    "--env",
                    # A throwaway key: only the loader is imported, no server runs.
                    f"WEBUI_SECRET_KEY={hashlib.sha256(run_dir.name.encode()).hexdigest()}",
                    "--env",
                    f"DOCLING_SERVER_URL={DOCLING_URL}",
                    "--env",
                    f"DOCLING_PARAMS={params}",
                    "--entrypoint",
                    "python3",
                    OPEN_WEBUI_IMAGE,
                    "-c",
                    _LOADER_PROGRAM,
                    "scan.pdf",
                    "text.pdf",
                ],
                capture_output=True,
                text=True,
                timeout=3600,
            )
            (run_dir / "loader.log").write_text(loaded.stdout + loaded.stderr)
            line = next(
                (row for row in loaded.stdout.splitlines() if row.startswith("LOADER-REPORT ")),
                None,
            )
            if line is None:
                print(
                    "Open WebUI's Docling loader returned no report (loader.log)", file=sys.stderr
                )
                return 1
            extracted = json.loads(line.removeprefix("LOADER-REPORT "))
        finally:
            subprocess.run(
                ["docker", "logs", container],
                stdout=(run_dir / "docling.log").open("w"),
                stderr=subprocess.STDOUT,
            )
            subprocess.run(["docker", "rm", "--force", container], capture_output=True)
            subprocess.run(["docker", "network", "rm", network], capture_output=True)
        scan = extracted["scan.pdf"]
        text = extracted["text.pdf"]
        flat_text = _normalized("".join(text["pages"]))
        report.update(
            scan_pages=len(scan["pages"]),
            scan_seconds=round(scan["seconds"], 1),
            scan_seconds_per_page=round(scan["seconds"] / pages, 1),
            scan_pages_with_key=[
                index + 1 for index, page in enumerate(scan["pages"]) if _has_key(page)
            ],
            text_seconds=round(text["seconds"], 1),
            text_missing_lines=[line for line in text_lines if _normalized(line) not in flat_text],
            scan_excerpt=scan["pages"][:1],
        )
        passed = (
            report["scan_pages"] == pages
            and report["scan_pages_with_key"] == list(range(1, pages + 1))
            and report["scan_seconds_per_page"] <= float(self.limits["seconds_per_page"])
            and not report["text_missing_lines"]
        )
        report["passed"] = passed
        (run_dir / "docling.json").write_text(
            json.dumps(report, ensure_ascii=False, indent=2) + "\n"
        )
        print(
            json.dumps({k: v for k, v in report.items() if k != "scan_excerpt"}, ensure_ascii=False)
        )
        return 0 if passed else 1

    def knowledge(self, run_dir: Path) -> int:
        """Japanese knowledge base in the Chat UI: OCR text, retrieval, ingestion and query time.

        ``KNOWLEDGE_KEEP=1`` keeps the knowledge base for the ``chat-ui`` gate;
        ``KNOWLEDGE_REUSE=1`` checks a kept one (after a restart) without ingesting.
        """
        token = self._token()
        existing = self._knowledge_by_name(token, KNOWLEDGE_NAME)
        report: dict = {"name": KNOWLEDGE_NAME, "reused": False}
        if existing is not None and os.environ.get("KNOWLEDGE_REUSE") == "1":
            knowledge_id = existing["id"]
            report["reused"] = True
            scan_id = next(
                row["id"]
                for row in self._knowledge_files(token, knowledge_id)
                if row["filename"] == KNOWLEDGE_SCAN
            )
        else:
            if existing is not None:
                self._delete_knowledge(token, existing["id"])
            knowledge_id = self._request(
                "POST",
                "/api/v1/knowledge/create",
                token,
                payload={
                    "name": KNOWLEDGE_NAME,
                    "description": "Kairyu gate: 試料の保存条件を含む日本語の技術報告書",
                },
            )["id"]
            scan = run_dir / KNOWLEDGE_SCAN
            self.scan_pdf(
                scan,
                int(self.limits["scan_pages"]),
                {int(self.limits["key_page"])},
                namespace=run_dir.name,
            )
            scan_id, seconds = self._ingest(
                token, knowledge_id, KNOWLEDGE_SCAN, scan.read_bytes(), "application/pdf"
            )
            report["scan_ingestion_seconds"] = round(seconds, 1)
            distractor_seconds = []
            for name, text in distractor_documents(
                int(self.limits["distractor_documents"]), namespace=run_dir.name
            ):
                distractor_seconds.append(
                    self._ingest(token, knowledge_id, name, text.encode(), "text/markdown")[1]
                )
            report["distractor_ingestion_seconds"] = round(sum(distractor_seconds), 1)
        content = (
            self._request("GET", f"/api/v1/files/{scan_id}/data/content", token).get("content")
            or ""
        )
        report["ocr_has_key"] = _has_key(content)
        report["ocr_characters"] = len(content)
        started = time.monotonic()
        result = self._request(
            "POST",
            "/api/v1/retrieval/query/collection",
            token,
            payload={"collection_names": [knowledge_id], "query": KNOWLEDGE_QUERY},
        )
        report["query_seconds"] = round(time.monotonic() - started, 2)
        documents = (result.get("documents") or [[]])[0]
        sources = [
            meta.get("name") or meta.get("source") for meta in (result.get("metadatas") or [[]])[0]
        ]
        top_n = int(self.limits["top_n"])
        report["returned"] = len(documents)
        report["key_rank"] = next(
            (rank + 1 for rank, text in enumerate(documents) if _has_key(text)), None
        )
        report["sources"] = sources
        report["scores"] = (result.get("distances") or [[]])[0]
        passed = (
            report["ocr_has_key"]
            and report["key_rank"] is not None
            and report["key_rank"] <= top_n
            and report["query_seconds"] <= float(self.limits["query_seconds"])
            and report.get("scan_ingestion_seconds", 0) <= float(self.limits["ingestion_seconds"])
        )
        report["passed"] = passed
        (run_dir / "knowledge.json").write_text(
            json.dumps(report, ensure_ascii=False, indent=2) + "\n"
        )
        print(json.dumps(report, ensure_ascii=False))
        if os.environ.get("KNOWLEDGE_KEEP") != "1":
            self._delete_knowledge(token, knowledge_id)
        return 0 if passed else 1

    def chat_ui(self, run_dir: Path) -> int:
        """The Chat UI in Chromium: 323, then 324 in one chat, Red for an image, a knowledge answer.

        The knowledge-base turn needs the base kept by ``KNOWLEDGE_KEEP=1 verify.sh knowledge``.
        """
        fixtures = run_dir / "fixtures"
        fixtures.mkdir(parents=True, exist_ok=True)
        from PIL import Image

        Image.new("RGB", (256, 256), (220, 20, 20)).save(fixtures / "red.png")
        before = self._kairyu_api_counts()
        browser = self._browser(
            run_dir,
            {
                "WEBUI_SMOKE_KNOWLEDGE": KNOWLEDGE_NAME,
                "WEBUI_SMOKE_KNOWLEDGE_QUERY": KNOWLEDGE_QUERY,
                "WEBUI_SMOKE_KNOWLEDGE_SOURCE": KNOWLEDGE_SCAN,
            },
            [f"{fixtures}:/fixtures:ro"],
        )
        after = self._kairyu_api_counts()
        report = {"before": before, "after": after, "browser_exit": browser.returncode}
        passed = (
            browser.returncode == 0
            and after["/v1/responses"] > before["/v1/responses"]
            and after["/v1/chat/completions"] == before["/v1/chat/completions"]
        )
        report["passed"] = passed
        (run_dir / "chat-ui.json").write_text(json.dumps(report, indent=2) + "\n")
        print(browser.stdout[-2000:] + browser.stderr[-2000:])
        print(json.dumps(report))
        token = self._token()
        kept = self._knowledge_by_name(token, KNOWLEDGE_NAME)
        if kept is not None and os.environ.get("KNOWLEDGE_KEEP") != "1":
            self._delete_knowledge(token, kept["id"])
        return 0 if passed else 1

    def knowledge_serving(self, run_dir: Path) -> int:
        """Serving rows while Docling OCRs scanned PDFs at its CPU cap: within the reference.

        Three load streams, one more than Docling's two conversion workers, keep
        a document queued at Docling while another stream embeds and adds its
        document; the second and third start when the first one's conversion
        ends. The rows start only once every stream has started and Docling is
        busy, and the gate fails unless Docling used at least
        ``docling_busy_cores`` in ``docling_busy_share`` of the seconds the rows
        took.
        """
        import queue

        limits = self.limits
        token = self._token()
        name = f"{KNOWLEDGE_NAME}-load"
        stale = self._knowledge_by_name(token, name)
        if stale is not None:
            self._delete_knowledge(token, stale["id"])
        knowledge_id = self._request(
            "POST",
            "/api/v1/knowledge/create",
            token,
            payload={"name": name, "description": "Kairyu gate load"},
        )["id"]
        # Open WebUI refuses a document whose content it already holds, so every
        # ingestion uploads a different scan; enough for the streams to outlast the rows.
        pending: queue.Queue[Path] = queue.Queue()
        for index in range(1, int(limits["load_documents"]) + 1):
            scan = run_dir / f"load-{index:02d}-{KNOWLEDGE_SCAN}"
            self.scan_pdf(
                scan,
                int(limits["scan_pages"]),
                {int(limits["key_page"])},
                namespace=f"{run_dir.name}-load-{index}",
            )
            pending.put(scan)
        stop = threading.Event()
        ingestions: list[dict] = []
        errors: list[str] = []

        def load(stream: int) -> None:
            while not stop.is_set():
                try:
                    scan = pending.get_nowait()
                except queue.Empty:
                    return
                started = time.monotonic()
                try:
                    self._ingest(
                        token, knowledge_id, scan.name, scan.read_bytes(), "application/pdf"
                    )
                except Exception as error:  # noqa: BLE001 - recorded and fails the gate
                    errors.append(f"stream {stream}: {error!r}")
                    stop.set()
                    return
                ingestions.append({"stream": stream, "start": started, "end": time.monotonic()})

        busy_cores = float(limits["docling_busy_cores"])
        concurrency = [int(level) for level in self.reference["concurrency"]]
        report: dict = {"load_streams": int(limits["load_streams"])}
        code = 1
        with DoclingCpu(f"{self.project}-docling-1") as cpu:
            workers = [
                threading.Thread(target=load, args=(stream,), daemon=True)
                for stream in range(1, int(limits["load_streams"]) + 1)
            ]
            # The first stream starts alone; the others start when its conversion
            # ends (Docling goes idle while it embeds and adds its document). With
            # one stream more than Docling's two workers, a document is always
            # queued at Docling while another stream embeds.
            busy_seconds = int(limits["docling_start_busy_seconds"])
            start_timeout = float(limits["docling_start_seconds"])
            workers[0].start()
            started_busy = cpu.wait_busy(busy_cores, busy_seconds, start_timeout)
            if len(workers) > 1:
                cpu.wait_idle(busy_cores, 2, float(limits["ingestion_seconds"]))
                for worker in workers[1:]:
                    worker.start()
                started_busy = started_busy and cpu.wait_busy(
                    busy_cores, busy_seconds, start_timeout
                )
            rows_started = time.monotonic()
            try:
                if started_busy and not errors:
                    code = self.serving(run_dir, concurrency)
            finally:
                rows_ended = time.monotonic()
                stop.set()
                for worker in workers:
                    worker.join(timeout=float(limits["ingestion_seconds"]) * 2)
                stuck = sum(worker.is_alive() for worker in workers)
            cpu.wait_past(rows_ended)
            load = cpu.summary(rows_started, rows_ended, busy_cores)
        if not stuck:
            self._delete_knowledge(token, knowledge_id)
        sampler_error = cpu.error
        rows = {}
        for level in concurrency:
            tok_s = self.serving_tok_s(run_dir, level)
            if tok_s is None:
                continue
            reference = float(self.reference["reference_tok_s"][str(level)])
            rows[str(level)] = {
                "output_tokens_per_s": round(tok_s, 1),
                "reference": reference,
                "fraction": round(tok_s / reference, 3),
            }
        report.update(
            docling_started_busy=started_busy,
            docling_sampler_error=sampler_error,
            docling_busy_share=round(load["busy_share"], 3),
            docling_observed_share=round(load["observed_share"], 3),
            docling_mean_cores=round(load["mean_cores"], 2),
            load_workers_stuck=stuck,
            rows_seconds=round(rows_ended - rows_started, 1),
            rows=rows,
            ingestions=[
                {
                    "stream": row["stream"],
                    "start_s": round(row["start"] - rows_started, 1),
                    "end_s": round(row["end"] - rows_started, 1),
                }
                for row in sorted(ingestions, key=lambda row: row["start"])
            ],
            errors=errors,
        )
        passed = (
            code == 0
            and started_busy
            and sampler_error is None
            and not errors
            and not stuck
            and load["busy_share"] >= float(limits["docling_busy_share"])
            and len(rows) == len(concurrency)
            and all(
                row["fraction"] >= float(self.reference["min_fraction"]) for row in rows.values()
            )
        )
        report["passed"] = passed
        (run_dir / "knowledge-serving.json").write_text(json.dumps(report, indent=2) + "\n")
        print(json.dumps(report))
        return 0 if passed else 1

    def office(self, run_dir: Path) -> int:
        """Word, Excel and PowerPoint in the Chat UI: read each attachment, then edit it to a file.

        Chromium attaches each fixture with the clip, asks about its key fact, then
        asks for an edit; the code interpreter writes the result to /mnt/uploads
        and the gate saves it from the reply's download link and checks it. Last,
        an image attached to that chat (after its tool calls) must be answered.
        """
        fixtures = run_dir / "fixtures"
        results = run_dir / "results"
        fixtures.mkdir(parents=True, exist_ok=True)
        results.mkdir(parents=True, exist_ok=True)
        os.chmod(results, 0o777)  # the browser container writes downloads here
        self._docling_python(_OFFICE_FIXTURES, fixtures, "/work")
        from PIL import Image

        Image.new("RGB", (256, 256), (220, 20, 20)).save(fixtures / "red.png")
        before = self._kairyu_api_counts()
        browser = self._browser(
            run_dir,
            {
                "WEBUI_SMOKE_MODE": "office",
                "WEBUI_SMOKE_OFFICE": json.dumps(OFFICE_CASES, ensure_ascii=False),
            },
            [f"{fixtures}:/fixtures:ro", f"{results}:/results"],
        )
        after = self._kairyu_api_counts()
        line = self._docling_python(_OFFICE_CHECK, results, "/work")
        checks = json.loads(line.split("OFFICE-CHECK ", 1)[1])
        report = {
            "browser_exit": browser.returncode,
            "checks": checks,
            "kairyu": {"before": before, "after": after},
        }
        passed = (
            browser.returncode == 0
            and all(row.get("passed") for row in checks.values())
            and after["/v1/responses"] > before["/v1/responses"]
            and after["/v1/chat/completions"] == before["/v1/chat/completions"]
        )
        report["passed"] = passed
        (run_dir / "office.json").write_text(
            json.dumps(report, ensure_ascii=False, indent=2) + "\n"
        )
        print(browser.stdout[-3000:] + browser.stderr[-2000:])
        print(json.dumps(report, ensure_ascii=False))
        return 0 if passed else 1

    def file_download(self, run_dir: Path) -> int:
        """A file the model makes is saved from the reply's link: desktop, after a reload, phone."""
        results = run_dir / "results"
        results.mkdir(parents=True, exist_ok=True)
        os.chmod(results, 0o777)  # the browser container writes downloads here
        before = self._kairyu_api_counts()
        browser = self._browser(
            run_dir,
            {"WEBUI_SMOKE_MODE": "download", "WEBUI_SMOKE_DOWNLOAD_REQUEST": DOWNLOAD_REQUEST},
            [f"{results}:/results"],
        )
        after = self._kairyu_api_counts()
        line = self._docling_python(_DOWNLOAD_CHECK, results, "/work")
        checks = json.loads(line.split("DOWNLOAD-CHECK ", 1)[1])
        report = {
            "browser_exit": browser.returncode,
            "checks": checks,
            "kairyu": {"before": before, "after": after},
        }
        passed = (
            browser.returncode == 0
            and all(
                checks[name].get("passed")
                for name in ("desktop.xlsx", "desktop-after-reload.xlsx", "phone.xlsx")
            )
            and checks["same_after_reload"]
            and after["/v1/responses"] > before["/v1/responses"]
            and after["/v1/chat/completions"] == before["/v1/chat/completions"]
        )
        report["passed"] = passed
        (run_dir / "file-download.json").write_text(
            json.dumps(report, ensure_ascii=False, indent=2) + "\n"
        )
        print(browser.stdout[-3000:] + browser.stderr[-2000:])
        print(json.dumps(report, ensure_ascii=False))
        return 0 if passed else 1
